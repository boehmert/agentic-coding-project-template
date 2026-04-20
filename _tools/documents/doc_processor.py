#!/usr/bin/env python3
"""
Dokument-Prozessor
Liest und konvertiert .docx, .pdf, .txt und .md Dateien.

Benötigte Pakete (in requirements.txt):
  python-docx  – Word-Dokumente
  pypdf        – PDF-Dokumente
  jinja2       – Template-Rendering
"""

from pathlib import Path


DEFAULT_REPORTS_DIR = Path("output/reports")


# ------------------------------------------------------------------ #
#  Word (.docx)                                                        #
# ------------------------------------------------------------------ #

def read_docx(file_path: Path) -> str:
    """
    Liest ein Word-Dokument (.docx) und gibt den Text zurück.

    Returns:
        Volltext als String, Absätze durch Zeilenumbrüche getrennt
    Raises:
        ImportError: Wenn python-docx nicht installiert ist
    """
    try:
        from docx import Document  # type: ignore
    except ImportError:
        raise ImportError("python-docx ist nicht installiert. Bitte 'pip install python-docx' ausführen.")

    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def _para_to_md(para) -> str:
    """Converts a single docx paragraph to a Markdown line."""
    text = para.text.strip()
    if not text:
        return ""
    style = para.style.name.lower() if para.style else ""
    if "heading 1" in style:
        return f"# {text}"
    elif "heading 2" in style:
        return f"## {text}"
    elif "heading 3" in style:
        return f"### {text}"
    elif "heading" in style:
        return f"#### {text}"
    else:
        return text


def _dedup_row(row: list[str]) -> list[str]:
    """Remove duplicate adjacent cell texts caused by merged cells in docx."""
    result = []
    prev = object()
    for cell in row:
        result.append(cell if cell != prev else "")
        prev = cell
    return result


def _kv_table_to_md(data: list[list[str]]) -> str:
    """Render a 2-column key/value table as a definition list."""
    lines = []
    for row in data:
        key, val = row[0], row[1]
        if key and val:
            lines.append(f"**{key}:** {val}")
        elif key:
            lines.append(f"**{key}**")
        elif val:
            lines.append(val)
    return "\n\n".join(lines)


def _fmt_row(cells: list[str]) -> str:
    """Format a list of cell texts as a Markdown table row."""
    return "| " + " | ".join(cells) + " |"


def _table_to_md(table) -> str:
    """
    Converts a docx table to a Markdown table.

    Strategy:
    - First row → header row
    - If the table looks like a 2-column key/value map (common in Jira exports),
      it is rendered as a definition list instead to avoid empty-cell clutter.
    """
    rows = table.rows
    if not rows:
        return ""

    # Collect all cell texts
    data = []
    for row in rows:
        data.append([cell.text.strip().replace("\n", " ") for cell in row.cells])

    if not data:
        return ""

    num_cols = len(data[0])
    data = [_dedup_row(row) for row in data]

    if num_cols == 2:
        return _kv_table_to_md(data)

    # General table → Markdown table
    header = data[0]
    separator = [":---" for _ in header]
    body = data[1:]

    lines = [_fmt_row(header), _fmt_row(separator)]
    for row in body:
        padded = (row + [""] * num_cols)[:num_cols]
        lines.append(_fmt_row(padded))
    return "\n".join(lines)


def _resolve_output_dir(output_dir: Path | None) -> Path:
    """Return the output directory and ensure it exists."""
    out = output_dir or DEFAULT_REPORTS_DIR
    out.mkdir(parents=True, exist_ok=True)
    return out


def _docx_block_to_md(child, paragraph_lookup: dict, table_lookup: dict) -> str:
    """Convert a body child element to Markdown if it is a paragraph or table."""
    tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
    if tag == "p":
        para = paragraph_lookup.get(child)
        return _para_to_md(para) if para is not None else ""
    if tag == "tbl":
        table = table_lookup.get(child)
        return _table_to_md(table) if table is not None else ""
    return ""


def _ordered_docx_blocks(doc) -> list[str]:
    """Render paragraphs and tables in document order."""
    paragraph_lookup = {para._element: para for para in doc.paragraphs}
    table_lookup = {tbl._element: tbl for tbl in doc.tables}

    blocks: list[str] = []
    for child in doc.element.body:
        md = _docx_block_to_md(child, paragraph_lookup, table_lookup)
        if md:
            blocks.append(md)
    return blocks


def _docx_output_stem(file_path: Path) -> str:
    """Strip a trailing .doc from names like foo.doc.docx."""
    stem = file_path.stem
    return stem[:-4] if stem.endswith(".doc") else stem


def docx_to_markdown(file_path: Path, output_dir: Path | None = None) -> Path:
    """
    Konvertiert ein Word-Dokument (.docx, auch .doc.docx) zu Markdown.

    Unterstützt:
    - Überschriften (Heading 1–4)
    - Fließtext-Absätze
    - Tabellen (2-Spalten → Definition List, mehrspaling → MD-Tabelle)
    - Reihenfolge: Absätze und Tabellen werden in Dokumentreihenfolge verarbeitet

    Hinweis: Für sehr komplexe Dokumente mit Bildern, Fußnoten oder
    verschachtelten Tabellen empfiehlt sich zusätzlich Pandoc (pypandoc).
    """
    try:
        from docx import Document  # type: ignore
    except ImportError:
        raise ImportError("python-docx ist nicht installiert.")

    doc = Document(file_path)
    out = _resolve_output_dir(output_dir)
    blocks = _ordered_docx_blocks(doc)
    stem = _docx_output_stem(file_path)

    md_content = "\n\n".join(blocks)
    out_file = out / (stem + ".md")
    out_file.write_text(md_content, encoding="utf-8")
    return out_file


# ------------------------------------------------------------------ #
#  PDF                                                                 #
# ------------------------------------------------------------------ #

def read_pdf(file_path: Path) -> str:
    """
    Liest ein PDF und gibt den Rohtext zurück.

    Raises:
        ImportError: Wenn pypdf nicht installiert ist
    """
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        raise ImportError("pypdf ist nicht installiert. Bitte 'pip install pypdf' ausführen.")

    reader = PdfReader(file_path)
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"<!-- Seite {i + 1} -->\n{text}")
    return "\n\n".join(pages)


def pdf_to_markdown(file_path: Path, output_dir: Path | None = None) -> Path:
    """Konvertiert ein PDF zu Markdown (roher Text)."""
    content = read_pdf(file_path)
    out = _resolve_output_dir(output_dir)
    out_file = out / (file_path.stem + ".md")
    out_file.write_text(f"# {file_path.stem}\n\n{content}", encoding="utf-8")
    return out_file


# ------------------------------------------------------------------ #
#  Markdown / Text                                                      #
# ------------------------------------------------------------------ #

def read_markdown(file_path: Path) -> str:
    """Liest eine Markdown-Datei"""
    return file_path.read_text(encoding="utf-8")


def read_text(file_path: Path) -> str:
    """Liest eine Textdatei"""
    return file_path.read_text(encoding="utf-8")


# ------------------------------------------------------------------ #
#  Jinja2-Template-Rendering                                           #
# ------------------------------------------------------------------ #

def render_template(template_str: str, context: dict) -> str:
    """
    Rendert einen Jinja2-Template-String mit dem gegebenen Kontext.

    Beispiel:
        render_template("# {{ title }}\\n\\n{{ content }}", {"title": "Mein Dokument", "content": "Text"})

    Raises:
        ImportError: Wenn jinja2 nicht installiert ist
    """
    try:
        from jinja2 import Template  # type: ignore
    except ImportError:
        raise ImportError("jinja2 ist nicht installiert. Bitte 'pip install jinja2' ausführen.")

    return Template(template_str).render(**context)


def render_template_file(template_path: Path, context: dict, output_path: Path) -> Path:
    """
    Rendert eine Jinja2-Template-Datei und speichert das Ergebnis.

    Args:
        template_path: Pfad zur Jinja2-Template-Datei
        context:       Dictionary mit Template-Variablen
        output_path:   Zieldatei

    Returns:
        Pfad zur erstellen Datei
    """
    try:
        from jinja2 import Environment, FileSystemLoader  # type: ignore
    except ImportError:
        raise ImportError("jinja2 ist nicht installiert.")

    env = Environment(
        loader=FileSystemLoader(str(template_path.parent)),
        autoescape=False,
    )
    template = env.get_template(template_path.name)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(template.render(**context), encoding="utf-8")
    return output_path


# ------------------------------------------------------------------ #
#  Batch-Verarbeitung                                                  #
# ------------------------------------------------------------------ #

def process_inbox(input_dir: Path = Path("inbox/docs"),
                  output_dir: Path = DEFAULT_REPORTS_DIR) -> list[Path]:
    """
    Verarbeitet alle unterstützten Dateien in inbox/docs/ und
    konvertiert sie zu Markdown nach output/reports/.

    Unterstützte Formate: .docx, .pdf, .txt, .md
    """
    if not input_dir.exists():
        print(f"⚠  Ordner nicht gefunden: {input_dir}")
        return []

    output_dir.mkdir(parents=True, exist_ok=True)
    results = []

    for file_path in sorted(input_dir.iterdir()):
        suffix = file_path.suffix.lower()
        try:
            if suffix == ".docx":
                out = docx_to_markdown(file_path, output_dir)
                print(f"✓ DOCX → MD: {file_path.name} → {out.name}")
                results.append(out)
            elif suffix == ".pdf":
                out = pdf_to_markdown(file_path, output_dir)
                print(f"✓ PDF  → MD: {file_path.name} → {out.name}")
                results.append(out)
            elif suffix in (".txt", ".md"):
                import shutil
                out = output_dir / file_path.name
                shutil.copy2(file_path, out)
                print(f"✓ Kopiert:   {file_path.name} → {out.name}")
                results.append(out)
            else:
                print(f"   Übersprungen (unbekanntes Format): {file_path.name}")
        except Exception as exc:
            print(f"✗ Fehler bei {file_path.name}: {exc}")

    return results


if __name__ == "__main__":
    process_inbox()

